#!/usr/bin/env python3
"""Excel 报表自动化 —— python-practice/02-excel-automation

任务卡：python-practice/02-excel-automation/README.md

这条流水线就是「把每周重复两小时的 Excel 手工活变成一个脚本」的标准形状：

    读多个 CSV  →  合并  →  清洗  →  统计聚合  →  写多 sheet 报表  →  画图

每一步都会打印自己做了什么、丢掉了多少行 —— 报表自动化最忌讳的就是
「悄悄地少了几行」。所有数字来自 data/ 下由 make_sample_data.py 生成的真实 CSV。

运行：
    python3 make_sample_data.py     # 先生成 data/*.csv
    python3 report.py               # 再跑报表
"""
from __future__ import annotations

import unicodedata
from pathlib import Path

import matplotlib

matplotlib.use("Agg")  # 无界面后端：服务器上没有显示器也能出图

import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
ASSETS = ROOT / "assets"

XLSX = ASSETS / "sales-report.xlsx"
CHART_MONTHLY = ASSETS / "chart-monthly-revenue.png"
CHART_CATEGORY = ASSETS / "chart-category-share.png"

# 配色沿用中国股市习惯：增长/正向用红，下降用绿
RED = "#c0392b"
GREEN = "#27ae60"

MONEY_COLS = ("销售额", "客单价", "收入", "金额", "revenue")


# ---------------------------------------------------------------------------
# 0. 环境准备：中文字体
# ---------------------------------------------------------------------------


def setup_chinese_font() -> str:
    """matplotlib 默认字体没有中文字形，不设就会画出一堆方框（豆腐块）。"""
    from matplotlib import font_manager

    installed = {f.name for f in font_manager.fontManager.ttflist}
    for name in ("PingFang SC", "Hiragino Sans GB", "Heiti SC", "Arial Unicode MS"):
        if name in installed:
            plt.rcParams["font.sans-serif"] = [name]
            plt.rcParams["axes.unicode_minus"] = False  # 负号不要变成方块
            return name
    return "(未找到中文字体，图中中文会显示为方块)"


# ---------------------------------------------------------------------------
# 1~2. 读取 + 合并 + 清洗
# ---------------------------------------------------------------------------


def load_and_clean() -> tuple[pd.DataFrame, dict[str, int]]:
    files = sorted(DATA.glob("sales-*.csv"))
    if not files:
        raise SystemExit("data/ 下没有 CSV，请先跑 `python3 make_sample_data.py`")

    frames = []
    for path in files:
        # utf-8-sig：写的时候带了 BOM，读的时候也要按带 BOM 解，否则第一列列名会多出 \ufeff
        frame = pd.read_csv(path, encoding="utf-8-sig")
        print(f"  读取 {path.name:<22} {len(frame):>5} 行")
        frames.append(frame)

    raw = pd.concat(frames, ignore_index=True)
    stats: dict[str, int] = {"合并后": len(raw)}

    df = raw.copy()

    # ① 日期：源数据混用 `2026-01-05` 与 `2026/01/05`，用 format="mixed" 让 pandas 自己判断
    df["date"] = pd.to_datetime(df["date"], format="mixed", errors="coerce")
    bad_date = int(df["date"].isna().sum())
    df = df.dropna(subset=["date"])

    # ② 去重：整行完全相同的记录只留一条
    dup = len(df) - len(df.drop_duplicates())
    df = df.drop_duplicates()

    # ③ 数量：空值无法参与计算，转成 NaN 后整行丢弃（真实业务里也可能改成填 0）
    df["qty"] = pd.to_numeric(df["qty"], errors="coerce")
    bad_qty = int(df["qty"].isna().sum())
    df = df.dropna(subset=["qty"])
    df["qty"] = df["qty"].astype(int)

    # ④ 派生列：销售额 = 数量 × 单价；月份列供透视表用
    df["revenue"] = (df["qty"] * df["unit_price"]).round(2)
    df["month"] = df["date"].dt.to_period("M").astype(str)

    stats.update(
        {"日期非法": bad_date, "重复行": dup, "数量为空": bad_qty, "清洗后": len(df)}
    )
    return df.sort_values("date").reset_index(drop=True), stats


# ---------------------------------------------------------------------------
# 3. 统计聚合
# ---------------------------------------------------------------------------


def aggregate(df: pd.DataFrame) -> dict[str, pd.DataFrame]:
    store = (
        df.groupby("store")
        .agg(
            订单数=("revenue", "size"),
            销售额=("revenue", "sum"),
            销量=("qty", "sum"),
            客单价=("revenue", "mean"),
        )
        .round(2)
        .sort_values("销售额", ascending=False)
    )
    store.index.name = "门店"

    category = (
        df.groupby("category")
        .agg(销售额=("revenue", "sum"), 销量=("qty", "sum"))
        .round(2)
        .sort_values("销售额", ascending=False)
    )
    category.index.name = "品类"
    category["占比"] = (category["销售额"] / category["销售额"].sum() * 100).round(1)

    pivot = df.pivot_table(
        index="month", columns="category", values="revenue", aggfunc="sum", margins=True
    ).round(2)
    pivot.index.name = "月份"
    pivot.columns.name = "品类"

    top_sku = (
        df.groupby(["sku", "category"])
        .agg(销售额=("revenue", "sum"), 销量=("qty", "sum"))
        .round(2)
        .sort_values("销售额", ascending=False)
        .head(10)
    )
    top_sku.index.names = ["SKU", "品类"]  # 样本里只有 9 个 SKU，所以表里就是全部排名

    monthly = (
        df.pivot_table(index="month", columns="store", values="revenue", aggfunc="sum")
        .round(2)
    )
    monthly.index.name = "月份"
    monthly.columns.name = "门店"

    return {
        "门店汇总": store,
        "品类汇总": category,
        "月份×品类": pivot,
        "SKU 排行": top_sku,
        "月份×门店": monthly,
    }


# ---------------------------------------------------------------------------
# 4. 写 Excel（多 sheet + 格式）
# ---------------------------------------------------------------------------


def disp_width(text: str) -> int:
    """按「终端显示宽度」算长度：中文/全角算 2，英文算 1。否则列宽会歪。"""
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in text)


def write_excel(df: pd.DataFrame, tables: dict[str, pd.DataFrame]) -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(XLSX, engine="openpyxl") as writer:
        df[
            ["date", "store", "category", "sku", "qty", "unit_price", "revenue"]
        ].to_excel(writer, sheet_name="清洗后明细", index=False)
        for name, table in tables.items():
            table.to_excel(writer, sheet_name=name, index=True)

    # 写完再用 openpyxl 打开做「人类友好」的加工 —— pandas 负责数据，openpyxl 负责颜面
    wb = load_workbook(XLSX)
    header_fill = PatternFill("solid", fgColor="2F5597")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"  # 冻结首行，滚动时表头不跑掉
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # 列宽按内容实际宽度自适应（中文字符按 2 格算）
        for col in ws.iter_cols():
            letter = col[0].column_letter
            widest = max(disp_width(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[letter].width = min(max(widest + 4, 10), 42)

        # 金额列加千分位与两位小数
        for idx, cell in enumerate(ws[1], start=1):
            if cell.value and any(k in str(cell.value) for k in MONEY_COLS):
                for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                    row[0].number_format = "#,##0.00"

        # 明细表的销售额列加「红高绿低」色阶，一眼看出大头在哪
        if ws.title == "清洗后明细":
            last = ws.max_row
            ws.conditional_formatting.add(
                f"G2:G{last}",
                ColorScaleRule(
                    start_type="min", start_color="E8F5E9",
                    mid_type="percentile", mid_value=50, mid_color="FFF9C4",
                    end_type="max", end_color="FFCDD2",
                ),
            )
    wb.save(XLSX)
    print(f"  写出 {XLSX.relative_to(ROOT.parent.parent)}  （{len(wb.worksheets)} 个 sheet）")


# ---------------------------------------------------------------------------
# 5. 画图
# ---------------------------------------------------------------------------


def draw_charts(tables: dict[str, pd.DataFrame]) -> None:
    monthly = tables["月份×门店"]
    ax = monthly.plot(
        kind="line", marker="o", linewidth=2.2, figsize=(9, 4.6), color=[RED, "#2F5597", GREEN]
    )
    ax.set_title("各门店月度销售额（2026 上半年）", fontsize=13, pad=12)
    ax.set_xlabel("月份")
    ax.set_ylabel("销售额（元）")
    ax.grid(alpha=0.3, linestyle="--")
    ax.legend(title="门店", frameon=False)
    # 保留一位小数：写成 .0f 的话 1.5 万会被舍成 2 万，刻度就全成了「2万」
    ax.yaxis.set_major_formatter(lambda v, _pos: f"{v / 10000:.1f}万")
    ax.figure.tight_layout()
    ax.figure.savefig(CHART_MONTHLY, dpi=130)
    plt.close(ax.figure)

    category = tables["品类汇总"]["销售额"]
    ax2 = category.plot(kind="barh", figsize=(8.4, 3.4), color="#2F5597", width=0.55)
    ax2.invert_yaxis()  # 销售额最高的排最上面
    ax2.set_title("品类销售额占比", fontsize=13, pad=12)
    ax2.set_xlabel("销售额（元）")
    ax2.set_ylabel("")
    ax2.grid(axis="x", alpha=0.3, linestyle="--")
    total = category.sum()
    for i, v in enumerate(category):
        ax2.text(v * 1.01, i, f"{v:,.0f}（{v / total * 100:.0f}%）", va="center", fontsize=9)
    ax2.set_xlim(0, category.max() * 1.28)
    ax2.figure.tight_layout()
    ax2.figure.savefig(CHART_CATEGORY, dpi=130)
    plt.close(ax2.figure)


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------


def main() -> int:
    font = setup_chinese_font()
    print(f"matplotlib 中文字体：{font}\n")

    print("① 读取三个门店的 CSV")
    df, stats = load_and_clean()

    print("\n② 合并与清洗")
    print(f"  合并后 {stats['合并后']} 行")
    print(f"  丢弃  重复行 {stats['重复行']} / 日期非法 {stats['日期非法']} / 数量为空 {stats['数量为空']}")
    print(f"  清洗后 {stats['清洗后']} 行（保留率 {stats['清洗后'] / stats['合并后'] * 100:.1f}%）")

    tables = aggregate(df)

    print("\n③ 门店汇总（按销售额降序）")
    print(tables["门店汇总"].to_string())

    print("\n④ 品类汇总")
    print(tables["品类汇总"].to_string())

    print("\n⑤ 月份 × 品类 透视表")
    print(tables["月份×品类"].to_string())

    print("\n⑥ SKU 销售额排行（样本共 9 个 SKU）")
    print(tables["SKU 排行"].to_string())

    print("\n⑦ 写 Excel + 画图")
    write_excel(df, tables)
    draw_charts(tables)
    for p in (CHART_MONTHLY, CHART_CATEGORY):
        print(f"  写出 {p.relative_to(ROOT.parent.parent)}")

    print("\n完成：一个原本要手工做两小时的周报，现在一条命令。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
